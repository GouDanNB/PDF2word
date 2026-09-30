"""Rule-based structure fixes for 应急预案-style DOCX after pdf2docx."""

from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from app.config import (
    BODY_FIRST_LINE_TWIPS,
    BODY_INDENT_FIX,
    BODY_INDENT_SYMMETRIC_MIN_TWIPS,
    BODY_LEFT_INDENT_MAX_TWIPS,
    HEADING_AUTO_NUMBER,
    HEADING_BOLD_CHAR_RATIO,
    HEADING_MAX_SIBLING_JUMP,
    HEADING_SIZE_ABOVE_BODY,
    HEADING_STRIP_PREFIX,
    HEADING_USE_PDF_OUTLINE,
)
from app.enhance.heading_numbering import apply_heading_number, ensure_heading_numbering
from app.enhance.pdf_outline import (
    OutlineHint,
    build_outline_index,
    extract_pdf_outline,
    match_outline_level,
)

# Soft line breaks / newlines that pdf2docx often leaves inside one <w:p>
_LINE_SPLIT_RE = re.compile(r"[\n\r\x0b]+")

# Split only when a real gap precedes a new heading (never inside 1.1.1)
_EMBEDDED_HEADING_RE = re.compile(
    r"(?<=\S)\s+(?="
    r"第[一二三四五六七八九十百零〇\d]+章"
    r"|\d+\.\d+(?:\.\d+)*(?:\s+|(?=[\u4e00-\u9fff]))"
    r")"
)

# Prefixes stripped when Word auto-numbering is enabled (must match detect rules)
_HEADING_PREFIX_RE = re.compile(
    r"^(?:"
    r"第[一二三四五六七八九十百零〇\d]+章\s*"
    r"|\d+\.\d+(?:\.\d+)?(?:\s+|(?=[\u4e00-\u9fffA-Za-z]))"
    r")"
)

# Reject TOC-like leaders: 第一章总则..........1
_TOC_LEADER_RE = re.compile(r"\.{3,}|…{2,}|…\s*\d+\s*$")

# Body list markers wrongly marked as Heading by older conversions / pdf2docx
_RE_PAREN_LIST_ITEM = re.compile(
    r"^[（(]\s*[一二三四五六七八九十百千零〇\d]+\s*[）)]"
)

# 第X章 / X.Y / X.Y.Z at line start
_RE_CHAPTER = re.compile(
    r"^第[一二三四五六七八九十百零〇\d]+章"
    r"(?:\s+|(?=[\u4e00-\u9fffA-Za-z]|$))"
)
_RE_XY = re.compile(
    r"^(\d+)\.(\d+)(?!\.\d)(?:\s+|(?=[\u4e00-\u9fffA-Za-z]|$))"
)
_RE_XYZ = re.compile(
    r"^(\d+)\.(\d+)\.(\d+)(?:\.\d+)*(?:\s+|(?=[\u4e00-\u9fffA-Za-z]|$))"
)
_RE_DOTTED_PREFIX = re.compile(r"^(\d+(?:\.\d+)+)")
_RE_CHAPTER_NUM = re.compile(r"^第([一二三四五六七八九十百零〇\d]+)章")

_CN_DIGITS = {
    "零": 0,
    "〇": 0,
    "一": 1,
    "二": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
    "十": 10,
}


class _OutlineNumberTracker:
    """Track accepted outline numbers; reject strong structural conflicts."""

    def __init__(self) -> None:
        self.chapter_idx = 0
        self.seen: set[tuple[int, ...]] = set()
        self.last_nums: tuple[int, ...] | None = None

    def note_chapter(self, idx: int) -> None:
        if idx > 0:
            self.chapter_idx = idx
        else:
            self.chapter_idx = max(self.chapter_idx, 0) + 1
        self.last_nums = (self.chapter_idx,)

    def note_numeric(self, nums: tuple[int, ...]) -> None:
        self.seen.add(nums)
        self.last_nums = nums
        if nums:
            self.chapter_idx = nums[0]

    def conflicts(self, nums: tuple[int, ...]) -> bool:
        """True if accepting nums would strongly contradict prior outline."""
        if not nums:
            return False
        if nums in self.seen:
            return True

        if len(nums) == 2:
            a, b = nums
            if self.chapter_idx > 0 and a != self.chapter_idx:
                # 正文里突然冒出与当前章不符的 X.Y
                return True
            if self.last_nums is not None:
                la = self.last_nums[0]
                if a < la:
                    return True
                if a > la + 1:
                    return True
                if a == la + 1 and b != 1:
                    return True
                if a == la and len(self.last_nums) >= 2:
                    lb = self.last_nums[1]
                    if b < lb:
                        return True
                    if b - lb > HEADING_MAX_SIBLING_JUMP:
                        return True
            return False

        if len(nums) >= 3:
            a, b, c = nums[0], nums[1], nums[2]
            parent = (a, b)
            if parent not in self.seen:
                # Allow if we just saw this parent as last H2
                if not (
                    self.last_nums is not None
                    and len(self.last_nums) >= 2
                    and self.last_nums[0] == a
                    and self.last_nums[1] == b
                ):
                    return True
            if self.chapter_idx > 0 and a != self.chapter_idx:
                return True
            if self.last_nums is not None and len(self.last_nums) >= 3:
                if self.last_nums[:2] == parent and c < self.last_nums[2]:
                    return True
                if self.last_nums[:2] == parent and c - self.last_nums[2] > HEADING_MAX_SIBLING_JUMP:
                    return True
            return False

        return False


class StructureFixer:
    """Promote headings, merge tiny fragments, light table cleanup."""

    def apply(
        self,
        docx_path: Path,
        *,
        pdf_path: Path | str | None = None,
        outline_hints: list[OutlineHint] | None = None,
    ) -> dict:
        path = Path(docx_path)
        doc = Document(str(path))
        stats = {
            "paragraphs_split": 0,
            "headings_promoted": 0,
            "headings_numbered": 0,
            "prefixes_stripped": 0,
            "fragments_merged": 0,
            "empty_removed": 0,
            "outline_hints": 0,
            "headings_outline_assisted": 0,
            "headings_outline_level_adjusted": 0,
            "headings_demoted": 0,
            "indents_normalized": 0,
        }

        hints = list(outline_hints or [])
        if (
            HEADING_USE_PDF_OUTLINE
            and not hints
            and pdf_path is not None
        ):
            try:
                hints = extract_pdf_outline(pdf_path)
            except Exception:  # noqa: BLE001 — outline is best-effort
                hints = []
        outline_index = build_outline_index(hints) if hints else {}
        stats["outline_hints"] = len(hints)

        # pdf2docx often packs title + body into one paragraph; split first
        # so Heading styles / outline levels apply per line.
        self._split_packed_paragraphs(doc, stats)

        num_id = None
        if HEADING_AUTO_NUMBER:
            try:
                num_id = ensure_heading_numbering(doc)
            except Exception:  # noqa: BLE001 — numbering is best-effort
                num_id = None

        self._promote_headings(
            doc, stats, num_id=num_id, outline_index=outline_index
        )
        # Drop leftover Heading styles on body list items (e.g. （1）预处理构筑物)
        self._demote_false_headings(doc, stats, outline_index=outline_index)
        self._merge_short_fragments(doc, stats)
        self._normalize_body_indents(doc, stats)
        self._normalize_tables(doc)

        doc.save(str(path))
        return stats

    def _split_packed_paragraphs(self, doc: Document, stats: dict) -> None:
        for para in list(doc.paragraphs):
            lines = self._extract_logical_lines(para.text or "")
            if len(lines) <= 1:
                continue

            # Capture formatting BEFORE wiping runs — critical for Chinese fonts
            # and partial-bold titles (e.g. only "1.1" bold, title text not).
            run_formats = self._collect_run_formats(para)
            default_r_pr = next(
                (r for _, r in run_formats if r is not None),
                None,
            )
            if default_r_pr is None and run_formats:
                default_r_pr = run_formats[0][1]
            p_pr_template = self._clone_p_pr(para)

            self._replace_paragraph_runs(
                para, self._runs_for_line(lines[0], run_formats, default_r_pr)
            )
            anchor = para
            for line in lines[1:]:
                segments = self._runs_for_line(line, run_formats, default_r_pr)
                anchor = self._insert_paragraph_after_runs(
                    anchor, segments, p_pr=p_pr_template
                )
            stats["paragraphs_split"] += len(lines) - 1

    def _extract_logical_lines(self, text: str) -> list[str]:
        raw = text.strip()
        if not raw:
            return []
        parts: list[str] = []
        for chunk in _LINE_SPLIT_RE.split(raw):
            chunk = chunk.strip()
            if not chunk:
                continue
            parts.extend(self._split_embedded_headings(chunk))
        return [p for p in parts if p]

    def _split_embedded_headings(self, text: str) -> list[str]:
        """Split '…正文 1.1 标题' into separate lines when safe."""
        pieces = _EMBEDDED_HEADING_RE.split(text)
        out: list[str] = []
        for piece in pieces:
            piece = piece.strip()
            if not piece:
                continue
            if out and self._detect_heading_level(piece) is None and len(piece) < 4:
                out[-1] = f"{out[-1]} {piece}".strip()
            else:
                out.append(piece)
        return out or [text]

    def _promote_headings(
        self,
        doc: Document,
        stats: dict,
        *,
        num_id: int | None,
        outline_index: dict[str, int] | None = None,
    ) -> None:
        tracker = _OutlineNumberTracker()
        body_median = self._estimate_body_median_size(doc)
        outline_index = outline_index or {}
        stats["body_median_pt"] = body_median
        stats["headings_rejected_unbold"] = 0
        stats["headings_rejected_conflict"] = 0

        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if not text or len(text) > 80:
                continue

            emphasized = self._para_is_emphasized(para, body_median)
            outline_lvl = (
                match_outline_level(text, outline_index) if outline_index else None
            )
            level = self._detect_heading_level(text, bold=emphasized)
            from_outline = False

            if level is None:
                pat = self._pattern_heading_level(text)
                # PDF bookmark/tag can rescue pattern titles that lost bold/size
                if pat is not None and outline_lvl is not None:
                    level = outline_lvl
                    from_outline = True
                    stats["headings_outline_assisted"] += 1
                elif pat is not None and pat >= 2 and not emphasized:
                    stats["headings_rejected_unbold"] += 1
                    continue
                else:
                    continue
            elif outline_lvl is not None and outline_lvl != level:
                level = outline_lvl
                from_outline = True
                stats["headings_outline_level_adjusted"] += 1

            nums = self._parse_dotted_nums(text)
            chapter_idx = (
                self._parse_chapter_index(text) if level == 1 else None
            )
            # Outline-confirmed titles skip structural conflict rejection
            if (
                level >= 2
                and nums is not None
                and not from_outline
                and outline_lvl is None
            ):
                if tracker.conflicts(nums):
                    stats["headings_rejected_conflict"] += 1
                    continue

            # Snapshot formatting BEFORE any text rewrite / style change
            saved_r_pr = self._first_content_r_pr(para)

            if HEADING_STRIP_PREFIX and HEADING_AUTO_NUMBER and num_id is not None:
                stripped = self._strip_heading_prefix(text)
                if stripped and stripped != text:
                    self._rewrite_para_keep_format(para, stripped)
                    stats["prefixes_stripped"] += 1
                    text = stripped
                if not text:
                    continue

            if self._apply_heading_style(para, level, saved_r_pr=saved_r_pr):
                stats["headings_promoted"] += 1
                if num_id is not None:
                    apply_heading_number(para, level, num_id)
                    stats["headings_numbered"] += 1
                if level == 1:
                    tracker.note_chapter(chapter_idx or 0)
                elif nums is not None:
                    tracker.note_numeric(nums)

    def _demote_false_headings(
        self,
        doc: Document,
        stats: dict,
        *,
        outline_index: dict[str, int] | None = None,
    ) -> None:
        """
        Demote paragraphs that still carry Heading/标题 style but are not
        real 第X章 / X.Y / X.Y.Z titles (common: （1）列表项 from old DOCX).
        """
        outline_index = outline_index or {}
        for para in doc.paragraphs:
            if self._is_body_style(para):
                continue
            text = (para.text or "").strip()
            if not text:
                continue

            # Keep pattern titles (before strip) and outline-confirmed wording
            if self._pattern_heading_level(text) is not None:
                continue
            if outline_index and match_outline_level(text, outline_index) is not None:
                continue

            # Demote clear false positives; keep short stripped titles (总则 etc.)
            if not self._is_false_heading_text(text):
                continue

            self._demote_to_body(para)
            stats["headings_demoted"] += 1

    @staticmethod
    def _is_false_heading_text(text: str) -> bool:
        t = (text or "").strip()
        if not t:
            return False
        # （1）预处理构筑物 / （一）指挥长职责
        if _RE_PAREN_LIST_ITEM.match(t):
            return True
        if _TOC_LEADER_RE.search(t):
            return True
        # Long body-like lines wrongly left as Heading
        if len(t) > 40:
            return True
        if re.search(r"[。；;]$", t):
            return True
        return False

    @staticmethod
    def _demote_to_body(para: Paragraph) -> None:
        """Remove Heading style / numbering / outline; keep run formatting."""
        try:
            para.style = "Normal"
        except (KeyError, ValueError):
            try:
                para.style = "正文"
            except (KeyError, ValueError):
                p_pr = para._p.get_or_add_pPr()
                p_style = p_pr.find(qn("w:pStyle"))
                if p_style is not None:
                    p_pr.remove(p_style)

        p_pr = para._p.get_or_add_pPr()
        for tag in ("w:numPr", "w:outlineLvl"):
            node = p_pr.find(qn(tag))
            if node is not None:
                p_pr.remove(node)

    @staticmethod
    def _strip_heading_prefix(text: str) -> str:
        stripped = _HEADING_PREFIX_RE.sub("", text, count=1).strip()
        # Keep original if stripping would leave nothing useful
        return stripped if stripped else text

    def _rewrite_para_keep_format(self, para: Paragraph, new_text: str) -> None:
        """
        Replace paragraph text, keeping title-body formatting.

        When only the number prefix was bold (pdf2docx partial bold), prefer the
        first run that overlaps the kept title wording — not the bold number run.
        """
        old = (para.text or "").strip()
        r_pr = self._r_pr_for_stripped_title(para, old, new_text)
        self._replace_paragraph_text(para, new_text, r_pr)

    @classmethod
    def _r_pr_for_stripped_title(
        cls, para: Paragraph, old_text: str, new_text: str
    ):
        """Pick rPr from the run covering title wording after prefix strip."""
        if not new_text:
            return cls._first_content_r_pr(para)

        # Prefer a run whose text appears in the kept title (not number-only).
        best = None
        best_score = -1
        for run in para.runs:
            t = (run.text or "").strip()
            if not t:
                continue
            existing = run._element.find(qn("w:rPr"))
            if existing is None:
                continue
            if t in new_text or new_text in t:
                score = min(len(t), len(new_text))
                # Prefer non-bold title runs when number was the only bold part
                if cls._run_is_bold(run) and t not in new_text:
                    score -= 100
                if score > best_score:
                    best_score = score
                    best = copy.deepcopy(existing)
            elif not cls._run_is_bold(run) and any(
                ch in t for ch in new_text[:2]
            ):
                # Weak fallback: unbold run near title start
                if best is None:
                    best = copy.deepcopy(existing)

        if best is not None:
            return best
        return cls._first_content_r_pr(para)

    def _apply_heading_style(
        self, para: Paragraph, level: int, *, saved_r_pr=None
    ) -> bool:
        """Set Word heading style + outlineLvl, keep original font/size/bold."""
        if saved_r_pr is None:
            saved_r_pr = self._first_content_r_pr(para)
        saved_runs = [
            (run.text, self._clone_element(run._element.find(qn("w:rPr"))))
            for run in para.runs
            if (run.text or "").strip()
        ]

        style_candidates = [
            f"Heading {level}",
            f"标题 {level}",
        ]
        applied = False
        for name in style_candidates:
            try:
                para.style = name
                applied = True
                break
            except (KeyError, ValueError):
                continue

        if not applied:
            for run in para.runs:
                run.bold = True
                run.font.size = Pt(18 - (level - 1) * 2)

        # Always restore PDF/docx run formatting (name/size/bold/eastAsia)
        fallback = saved_r_pr or next(
            (r for _, r in saved_runs if r is not None), None
        )
        for run in para.runs:
            if not (run.text or "").strip():
                continue
            matched = False
            for text, r_pr in saved_runs:
                if run.text == text and r_pr is not None:
                    self._set_run_r_pr(run, r_pr)
                    matched = True
                    break
            if not matched and fallback is not None:
                self._set_run_r_pr(run, fallback)

        self._set_outline_level(para, level)
        return True

    @staticmethod
    def _set_outline_level(para: Paragraph, level: int) -> None:
        """Word outlineLvl is 0-based (Heading 1 → 0)."""
        p_pr = para._p.get_or_add_pPr()
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            p_pr.append(outline)
        outline.set(qn("w:val"), str(max(0, min(level, 9) - 1)))

    def _detect_heading_level(
        self,
        text: str,
        *,
        size_pt: float | None = None,
        bold: bool = False,
        font_name: str = "",
        body_median: float | None = None,
    ) -> int | None:
        """
        Only explicit patterns become headings:
          第X章 → 1
          X.Y / X.Y.Z → 2/3 only if emphasized (real bold majority or larger size)
        """
        del size_pt, font_name, body_median  # call-site compatibility
        text = (text or "").strip()
        if not text or len(text) > 80:
            return None
        if _TOC_LEADER_RE.search(text):
            return None

        pattern_level = self._pattern_heading_level(text)
        if pattern_level is None:
            return None
        # 数字式标题：须「强调」(多数字符加粗，或明显大于正文字号)
        if pattern_level >= 2 and not bold:
            return None
        return pattern_level

    @staticmethod
    def _pattern_heading_level(text: str) -> int | None:
        if _RE_CHAPTER.match(text):
            return 1
        if _RE_XYZ.match(text):
            return 3
        if _RE_XY.match(text):
            return 2
        return None

    @staticmethod
    def _parse_dotted_nums(text: str) -> tuple[int, ...] | None:
        m = _RE_DOTTED_PREFIX.match((text or "").strip())
        if not m:
            return None
        return tuple(int(p) for p in m.group(1).split("."))

    @staticmethod
    def _parse_chapter_index(text: str) -> int | None:
        """Parse 第一章 / 第1章 → 1. Returns None if unparsable."""
        m = _RE_CHAPTER_NUM.match((text or "").strip())
        if not m:
            return None
        token = m.group(1)
        if token.isdigit():
            return int(token)
        if token == "十":
            return 10
        if token.startswith("十"):
            return 10 + _CN_DIGITS.get(token[1:], 0)
        if "十" in token:
            left, _, right = token.partition("十")
            return _CN_DIGITS.get(left, 0) * 10 + _CN_DIGITS.get(right, 0)
        return _CN_DIGITS.get(token)

    @classmethod
    def _para_is_emphasized(cls, para: Paragraph, body_median: float | None) -> bool:
        """
        True if the paragraph is visually emphasized like a heading.

        pdf2docx often bolds only the number (``1.1``) while title text is not;
        ``w:b val=\"0\"`` is ignored. Signals (any one is enough):

        - Leading chapter/number prefix is truly bold
        - Character-weighted bold ratio >= HEADING_BOLD_CHAR_RATIO
        - Typical/majority run size clearly larger than body median
        """
        bold_chars = 0
        total_chars = 0
        sizes: list[float] = []
        for run in para.runs:
            raw = run.text or ""
            t = raw.strip()
            if not t:
                continue
            n = len(t)
            total_chars += n
            if cls._run_is_bold(run):
                bold_chars += n
            sz = cls._run_size_pt(run)
            if sz is not None:
                sizes.extend([sz] * n)

        if cls._leading_number_is_bold(para):
            return True

        if total_chars and (bold_chars / total_chars) >= HEADING_BOLD_CHAR_RATIO:
            return True

        if sizes and body_median is not None:
            sizes_sorted = sorted(sizes)
            typical = sizes_sorted[len(sizes_sorted) // 2]
            if typical >= body_median + HEADING_SIZE_ABOVE_BODY:
                return True
            if max(sizes) >= body_median + HEADING_SIZE_ABOVE_BODY + 0.5:
                large = sum(
                    1 for s in sizes if s >= body_median + HEADING_SIZE_ABOVE_BODY
                )
                if large / len(sizes) >= 0.5:
                    return True
        return False

    @classmethod
    def _leading_number_is_bold(cls, para: Paragraph) -> bool:
        """True when the line-leading 第X章 / X.Y / X.Y.Z span is truly bold."""
        text = (para.text or "").strip()
        if not text:
            return False
        m = _RE_CHAPTER_NUM.match(text) or _RE_DOTTED_PREFIX.match(text)
        if not m:
            return False
        prefix = m.group(0)
        # Walk runs covering the prefix characters
        pos = 0
        covered_bold = 0
        covered_total = 0
        for run in para.runs:
            raw = run.text or ""
            if not raw:
                continue
            for ch in raw:
                if pos >= len(prefix):
                    break
                # Skip leading whitespace in paragraph text already stripped
                if ch.isspace() and covered_total == 0:
                    continue
                if pos < len(prefix) and ch == prefix[pos]:
                    covered_total += 1
                    if cls._run_is_bold(run):
                        covered_bold += 1
                    pos += 1
                elif covered_total == 0:
                    continue
                else:
                    break
            if pos >= len(prefix):
                break
        if covered_total < max(1, len(prefix) // 2):
            return False
        return covered_bold >= max(1, (covered_total + 1) // 2)

    @classmethod
    def _para_font_metrics(cls, para: Paragraph) -> tuple[float | None, bool, str]:
        """Dominant size (pt), emphasized?, first eastAsia/ascii font name."""
        sizes: list[float] = []
        font_name = ""
        for run in para.runs:
            if not (run.text or "").strip():
                continue
            sz = cls._run_size_pt(run)
            if sz is not None:
                sizes.append(sz)
            if not font_name:
                font_name = cls._run_font_name(run)
        size_pt = max(sizes) if sizes else None
        return size_pt, cls._para_is_emphasized(para, None), font_name

    @classmethod
    def _estimate_body_median_size(cls, doc: Document) -> float | None:
        """Median font size of longer non-pattern paragraphs (≈ body)."""
        sizes: list[float] = []
        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if len(text) < 25:
                continue
            if cls._pattern_heading_level(text) is not None:
                continue
            for run in para.runs:
                if not (run.text or "").strip():
                    continue
                sz = cls._run_size_pt(run)
                if sz is not None:
                    sizes.append(sz)
        if not sizes:
            for para in doc.paragraphs:
                for run in para.runs:
                    if not (run.text or "").strip():
                        continue
                    sz = cls._run_size_pt(run)
                    if sz is not None:
                        sizes.append(sz)
        if not sizes:
            return None
        sizes.sort()
        return sizes[len(sizes) // 2]

    @staticmethod
    def _run_is_bold(run) -> bool:
        """
        True only for real bold.
        Explicit w:b w:val=\"0\" / false / off → not bold (common from pdf2docx).
        """
        r_pr = run._element.find(qn("w:rPr"))
        if r_pr is not None:
            for tag in ("w:b", "w:bCs"):
                el = r_pr.find(qn(tag))
                if el is None:
                    continue
                val = el.get(qn("w:val"))
                if val is None:
                    return True
                if val in ("0", "false", "off"):
                    return False
                return True
        # Fall back to python-docx (may be None when inherited)
        if run.bold is True:
            return True
        if run.bold is False:
            return False
        # Font name hints only when name clearly includes Bold/Black
        name = (run.font.name or "").lower()
        if re.search(r"(?:^|[-_ ])(bold|black|heavy)(?:$|[-_ ])", name):
            return True
        return False

    @staticmethod
    def _run_size_pt(run) -> float | None:
        if run.font.size is not None:
            return float(run.font.size.pt)
        r_pr = run._element.find(qn("w:rPr"))
        if r_pr is None:
            return None
        sz = r_pr.find(qn("w:sz"))
        if sz is not None and sz.get(qn("w:val")):
            return int(sz.get(qn("w:val"))) / 2.0
        sz_cs = r_pr.find(qn("w:szCs"))
        if sz_cs is not None and sz_cs.get(qn("w:val")):
            return int(sz_cs.get(qn("w:val"))) / 2.0
        return None

    @staticmethod
    def _run_font_name(run) -> str:
        r_pr = run._element.find(qn("w:rPr"))
        if r_pr is not None:
            r_fonts = r_pr.find(qn("w:rFonts"))
            if r_fonts is not None:
                for attr in ("w:eastAsia", "w:ascii", "w:hAnsi", "w:cs"):
                    val = r_fonts.get(qn(attr))
                    if val:
                        return val
        return run.font.name or ""

    @classmethod
    def _first_content_r_pr(cls, para: Paragraph):
        for run in para.runs:
            if (run.text or "").strip():
                existing = run._element.find(qn("w:rPr"))
                if existing is not None:
                    return copy.deepcopy(existing)
        return None

    def _normalize_body_indents(self, doc: Document, stats: dict) -> None:
        """
        Fix pdf2docx body indents: left-only (全体缩进) → firstLine (首行缩进).

        Skips headings, TOC leaders, hanging list indents, and symmetric
        cover/centered blocks (left≈right).
        """
        if not BODY_INDENT_FIX:
            return
        for para in doc.paragraphs:
            if self._fix_body_indent(para):
                stats["indents_normalized"] += 1
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if self._fix_body_indent(para):
                            stats["indents_normalized"] += 1

    @classmethod
    def _fix_body_indent(cls, para: Paragraph) -> bool:
        if not cls._is_body_style(para):
            return False
        text = (para.text or "").strip()
        if not text or _TOC_LEADER_RE.search(text):
            return False

        p_pr = para._p.find(qn("w:pPr"))
        if p_pr is None:
            return False

        jc = p_pr.find(qn("w:jc"))
        if jc is not None and jc.get(qn("w:val")) in ("center", "right", "both"):
            # both/distribute kept; center/right covers skip
            if jc.get(qn("w:val")) in ("center", "right"):
                return False

        ind = p_pr.find(qn("w:ind"))
        if ind is None:
            return False

        left = cls._ind_attr_int(ind, "w:left", "w:start")
        right = cls._ind_attr_int(ind, "w:right", "w:end")
        first = cls._ind_attr_int(ind, "w:firstLine")
        hanging = cls._ind_attr_int(ind, "w:hanging")

        # Hanging list layout — leave alone
        if hanging > 0:
            return False

        # Cover / centered block: large symmetric left+right
        if (
            left >= BODY_INDENT_SYMMETRIC_MIN_TWIPS
            and right >= BODY_INDENT_SYMMETRIC_MIN_TWIPS
            and abs(left - right) <= 120
        ):
            return False

        changed = False

        # Double-count: left + firstLine → drop left, keep firstLine
        if left > 0 and first > 0:
            ind.attrib.pop(qn("w:left"), None)
            ind.attrib.pop(qn("w:start"), None)
            changed = True
            left = 0

        # Left-only body indent → first-line indent
        if left > 0 and first == 0 and left <= BODY_LEFT_INDENT_MAX_TWIPS:
            # Prefer standard 2-char indent when left is in typical body range
            fl = (
                BODY_FIRST_LINE_TWIPS
                if 200 <= left <= 600
                else left
            )
            ind.set(qn("w:firstLine"), str(fl))
            ind.attrib.pop(qn("w:left"), None)
            ind.attrib.pop(qn("w:start"), None)
            # Drop tiny right indent that often accompanies pdf2docx body
            if 0 < right <= 300:
                ind.attrib.pop(qn("w:right"), None)
                ind.attrib.pop(qn("w:end"), None)
            changed = True

        # Remove empty <w:ind/>
        if changed and len(ind.attrib) == 0:
            p_pr.remove(ind)

        return changed

    @staticmethod
    def _ind_attr_int(ind, *attrs: str) -> int:
        for attr in attrs:
            val = ind.get(qn(attr))
            if val is not None:
                try:
                    return int(val)
                except ValueError:
                    return 0
        return 0

    def _merge_short_fragments(self, doc: Document, stats: dict) -> None:
        """Merge consecutive 1–2 character body paragraphs into previous."""
        paragraphs = list(doc.paragraphs)
        i = 1
        while i < len(paragraphs):
            prev = paragraphs[i - 1]
            cur = paragraphs[i]
            prev_t = (prev.text or "").strip()
            cur_t = (cur.text or "").strip()
            if (
                cur_t
                and len(cur_t) <= 2
                and prev_t
                and not self._looks_like_heading(prev_t)
                and not self._looks_like_heading(cur_t)
                and self._is_body_style(cur)
                and self._is_body_style(prev)
            ):
                if prev.runs:
                    prev.runs[-1].text = (prev.runs[-1].text or "") + cur_t
                else:
                    prev.add_run(cur_t)
                self._clear_paragraph(cur)
                stats["fragments_merged"] += 1
            i += 1

    def _normalize_tables(self, doc: Document) -> None:
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if run.text:
                                run.text = run.text.replace("\u00a0", " ").strip()

    @classmethod
    def _insert_paragraph_after(
        cls,
        paragraph: Paragraph,
        text: str,
        *,
        r_pr=None,
        p_pr=None,
    ) -> Paragraph:
        return cls._insert_paragraph_after_runs(
            paragraph, [(text, r_pr)] if text else [], p_pr=p_pr
        )

    @classmethod
    def _insert_paragraph_after_runs(
        cls,
        paragraph: Paragraph,
        segments: list[tuple[str, object]],
        *,
        p_pr=None,
    ) -> Paragraph:
        new_p = OxmlElement("w:p")
        if p_pr is not None:
            new_p.insert(0, copy.deepcopy(p_pr))
        paragraph._p.addnext(new_p)
        new_para = Paragraph(new_p, paragraph._parent)
        for text, r_pr in segments:
            if not text:
                continue
            run = new_para.add_run(text)
            if r_pr is not None:
                cls._set_run_r_pr(run, r_pr)
        return new_para

    @classmethod
    def _replace_paragraph_text(cls, para: Paragraph, text: str, r_pr=None) -> None:
        cls._replace_paragraph_runs(para, [(text, r_pr)] if text else [])

    @classmethod
    def _replace_paragraph_runs(
        cls, para: Paragraph, segments: list[tuple[str, object]]
    ) -> None:
        p = para._p
        for child in list(p):
            if child.tag == qn("w:r"):
                p.remove(child)
        for text, r_pr in segments:
            if not text:
                continue
            run = para.add_run(text)
            if r_pr is not None:
                cls._set_run_r_pr(run, r_pr)

    @staticmethod
    def _collect_run_formats(para: Paragraph) -> list[tuple[str, object]]:
        """[(run_text, cloned rPr or None), ...] — keep newlines for line mapping."""
        out: list[tuple[str, object]] = []
        for run in para.runs:
            text = run.text or ""
            if text == "":
                continue
            r_pr = run._element.find(qn("w:rPr"))
            out.append((text, copy.deepcopy(r_pr) if r_pr is not None else None))
        return out

    @classmethod
    def _runs_for_line(
        cls,
        line: str,
        run_formats: list[tuple[str, object]],
        default_r_pr,
    ) -> list[tuple[str, object]]:
        """
        Map a logical line back onto original runs so partial bold is preserved.

        Example: runs [bold\"1.1\", plain\"编制目的\"] stay two runs after split.
        """
        line = (line or "").strip()
        if not line:
            return []
        if not run_formats:
            return [(line, default_r_pr)]

        # Build concatenated plain text of runs (strip-aware: use raw texts)
        pieces: list[tuple[str, object]] = []
        for text, r_pr in run_formats:
            if text is None:
                continue
            pieces.append((text, r_pr))
        full = "".join(t for t, _ in pieces)
        # Locate line inside full (tolerate whitespace differences)
        idx = full.find(line)
        if idx < 0:
            compact_full = re.sub(r"\s+", "", full)
            compact_line = re.sub(r"\s+", "", line)
            idx_c = compact_full.find(compact_line)
            if idx_c < 0:
                return [(line, cls._r_pr_for_line(line, run_formats, default_r_pr))]
            # Fall back to single preferred rPr when whitespace diverges
            return [(line, cls._r_pr_for_line(line, run_formats, default_r_pr))]

        end = idx + len(line)
        segments: list[tuple[str, object]] = []
        cursor = 0
        for text, r_pr in pieces:
            run_start = cursor
            run_end = cursor + len(text)
            cursor = run_end
            if run_end <= idx or run_start >= end:
                continue
            slice_start = max(run_start, idx) - run_start
            slice_end = min(run_end, end) - run_start
            frag = text[slice_start:slice_end]
            if frag:
                segments.append((frag, r_pr if r_pr is not None else default_r_pr))

        if not segments:
            return [(line, cls._r_pr_for_line(line, run_formats, default_r_pr))]
        # If line had surrounding strip difference, ensure joined text matches
        joined = "".join(t for t, _ in segments)
        if joined.strip() != line and joined.replace("\n", "").strip() != line:
            return [(line, cls._r_pr_for_line(line, run_formats, default_r_pr))]
        return segments

    @staticmethod
    def _r_pr_is_bold(r_pr) -> bool:
        if r_pr is None:
            return False
        for tag in ("w:b", "w:bCs"):
            el = r_pr.find(qn(tag))
            if el is None:
                continue
            val = el.get(qn("w:val"))
            if val is None:
                return True
            if val in ("0", "false", "off"):
                return False
            return True
        return False

    @staticmethod
    def _r_pr_size_half_points(r_pr) -> int | None:
        if r_pr is None:
            return None
        sz = r_pr.find(qn("w:sz"))
        if sz is not None and sz.get(qn("w:val")):
            try:
                return int(sz.get(qn("w:val")))
            except ValueError:
                return None
        return None

    @classmethod
    def _r_pr_for_line(
        cls,
        line: str,
        run_formats: list[tuple[str, object]],
        default_r_pr,
    ):
        """
        Pick one run format for a logical line (fallback when multi-run map fails).

        Prefer larger size; among equals, prefer non-bold title-body runs over a
        short bold number-only run so we don't paint the whole title bold.
        """
        if not run_formats:
            return default_r_pr
        candidates: list[tuple[int, int, int, object]] = []
        for text, r_pr in run_formats:
            if not text or r_pr is None:
                continue
            stripped = text.strip()
            if not (line == stripped or line in text or stripped in line):
                continue
            overlap = min(len(line), len(stripped))
            size = cls._r_pr_size_half_points(r_pr) or 0
            # Score: size first, then overlap; penalize number-only bold fragments
            number_only = bool(re.fullmatch(r"\d+(?:\.\d+)*", stripped))
            bold = cls._r_pr_is_bold(r_pr)
            penalty = 0
            if number_only and bold and len(stripped) < len(line):
                penalty = 50
            candidates.append((size, overlap - penalty, overlap, r_pr))
        if not candidates:
            return default_r_pr
        candidates.sort(key=lambda x: (x[0], x[1], x[2]), reverse=True)
        return candidates[0][3]

    @staticmethod
    def _clone_p_pr(para: Paragraph):
        """Clone paragraph properties but drop numbering / outline (set later)."""
        p_pr = para._p.find(qn("w:pPr"))
        if p_pr is None:
            return None
        cloned = copy.deepcopy(p_pr)
        for tag in ("w:numPr", "w:outlineLvl", "w:pStyle"):
            node = cloned.find(qn(tag))
            if node is not None:
                cloned.remove(node)
        return cloned

    @staticmethod
    def _clone_element(el):
        return copy.deepcopy(el) if el is not None else None

    @staticmethod
    def _set_run_r_pr(run, r_pr) -> None:
        if r_pr is None:
            return
        r = run._element
        existing = r.find(qn("w:rPr"))
        if existing is not None:
            r.remove(existing)
        r.insert(0, copy.deepcopy(r_pr))

    @staticmethod
    def _looks_like_heading(text: str) -> bool:
        t = (text or "").strip()
        if not t or _TOC_LEADER_RE.search(t):
            return False
        return StructureFixer._pattern_heading_level(t) is not None

    @staticmethod
    def _is_body_style(para) -> bool:
        name = (para.style.name if para.style else "") or ""
        return not (
            name.startswith("Heading")
            or name.startswith("标题")
        )

    @staticmethod
    def _clear_paragraph(para) -> None:
        for run in para.runs:
            run.text = ""
