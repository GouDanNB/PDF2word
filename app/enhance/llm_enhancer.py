"""Apply LLM structure instructions to DOCX without rewriting text."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document

from app.enhance.llm_client import LLMClient


class LLMEnhancer:
    """Extract heading candidates → LLM classify → apply styles only."""

    def __init__(self, client: LLMClient | None = None) -> None:
        self.client = client or LLMClient()

    def enhance(self, docx_path: Path) -> dict[str, Any]:
        path = Path(docx_path)
        if not self.client.is_available():
            return {
                "applied": 0,
                "skipped": True,
                "reason": "本地 Ollama 未运行，已跳过 LLM 增强",
            }

        doc = Document(str(path))
        candidates: list[dict[str, Any]] = []
        for i, para in enumerate(doc.paragraphs):
            text = (para.text or "").strip()
            if not text or len(text) > 80:
                continue
            # Prefer short lines that look structural
            if len(text) <= 60:
                candidates.append({"index": i, "text": text})

        if not candidates:
            return {"applied": 0, "skipped": False, "reason": "无候选标题"}

        try:
            items = self.client.classify_headings(candidates)
        except Exception as exc:  # noqa: BLE001 — surface as soft failure
            return {
                "applied": 0,
                "skipped": True,
                "reason": f"LLM 调用失败，已跳过: {exc}",
            }

        applied = 0
        by_index = {it["index"]: it["style"] for it in items}
        for i, para in enumerate(doc.paragraphs):
            style = by_index.get(i)
            if not style or style == "Body":
                continue
            try:
                para.style = style
                applied += 1
            except KeyError:
                continue

        doc.save(str(path))
        return {"applied": applied, "skipped": False, "reason": "ok"}
