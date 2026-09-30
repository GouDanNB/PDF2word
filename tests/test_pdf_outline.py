"""Tests for PDF outline extraction and heading assist."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Pt

from app.enhance.pdf_outline import (
    OutlineHint,
    build_outline_index,
    extract_pdf_outline,
    match_outline_level,
    normalize_outline_title,
)
from app.enhance.structure_fix import StructureFixer


class PdfOutlineTests(unittest.TestCase):
    def test_normalize_and_match(self) -> None:
        self.assertEqual(normalize_outline_title("1.1 编制目的"), "1.1编制目的")
        self.assertEqual(normalize_outline_title("第一章  总则"), "第一章总则")
        index = build_outline_index(
            [
                OutlineHint(1, "第一章 总则", source="toc"),
                OutlineHint(2, "1.1 编制目的", source="toc"),
                OutlineHint(3, "1.2.1 有关法律法规和要求", source="toc"),
            ]
        )
        self.assertEqual(match_outline_level("第一章总则", index), 1)
        self.assertEqual(match_outline_level("1.1编制目的", index), 2)
        self.assertEqual(match_outline_level("编制目的", index), 2)
        self.assertEqual(match_outline_level("1.2.1有关法律法规和要求", index), 3)
        self.assertIsNone(match_outline_level("普通正文段落", index))

    def test_extract_from_houzhai_pdf_if_present(self) -> None:
        pdf = Path(r"d:\我的软件\李宇轩\测试\修编\15、后宅运营部应急预案.pdf")
        if not pdf.is_file():
            self.skipTest("sample PDF not available")
        hints = extract_pdf_outline(pdf)
        self.assertGreater(len(hints), 50)
        self.assertTrue(any(h.level == 1 and "总则" in h.title for h in hints))
        self.assertTrue(any(h.level == 2 and "编制目的" in h.title for h in hints))
        self.assertTrue(all(h.source == "toc" for h in hints))


class OutlineAssistStructureTests(unittest.TestCase):
    def test_outline_rescues_unemphasized_numeric_title(self) -> None:
        """Bookmark match promotes X.Y even when bold/size emphasis is missing."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "t.docx"
            doc = Document()
            body = doc.add_paragraph()
            br = body.add_run("这是一段足够长的正文内容用来估计正文字号中位数十二点。")
            br.font.size = Pt(12)
            p0 = doc.add_paragraph()
            r0 = p0.add_run("第一章 总则")
            r0.font.size = Pt(12)  # same as body, no bold
            p1 = doc.add_paragraph()
            r1 = p1.add_run("1.1 编制目的")
            r1.font.size = Pt(12)
            r1.bold = False
            p2 = doc.add_paragraph()
            r2 = p2.add_run("1.9 正文里的引用编号不应升")
            r2.font.size = Pt(12)
            doc.save(str(path))

            hints = [
                OutlineHint(1, "第一章 总则", source="toc"),
                OutlineHint(2, "1.1 编制目的", source="toc"),
            ]
            stats = StructureFixer().apply(path, outline_hints=hints)
            self.assertGreaterEqual(stats["outline_hints"], 2)
            self.assertGreaterEqual(stats["headings_outline_assisted"], 1)
            self.assertGreaterEqual(stats["headings_promoted"], 2)

            doc2 = Document(str(path))
            by_style = {
                (p.text or "").strip(): (p.style.name if p.style else "")
                for p in doc2.paragraphs
                if (p.text or "").strip()
            }
            self.assertTrue(by_style.get("总则", "").startswith("Heading"))
            self.assertTrue(by_style.get("编制目的", "").startswith("Heading"))
            # Not in TOC → stays body when unemphasized
            self.assertFalse(
                (by_style.get("1.9 正文里的引用编号不应升", "") or "").startswith(
                    "Heading"
                )
            )


if __name__ == "__main__":
    unittest.main()
