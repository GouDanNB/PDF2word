"""Unit tests for heading heuristics."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from app.enhance.structure_fix import StructureFixer
from app.enhance.llm_client import LLMClient


class StructureFixerTests(unittest.TestCase):
    def test_detect_levels(self) -> None:
        fixer = StructureFixer()
        # 第X章: no bold required
        self.assertEqual(fixer._detect_heading_level("第一章 总则"), 1)
        self.assertEqual(fixer._detect_heading_level("第1章总则"), 1)
        # Numeric: must be bold (heading font alone is NOT enough)
        self.assertEqual(
            fixer._detect_heading_level("1.1 编制目的", bold=True), 2
        )
        self.assertEqual(
            fixer._detect_heading_level("1.1编制目的", bold=True), 2
        )
        self.assertEqual(
            fixer._detect_heading_level("1.1.1 细则", bold=True), 3
        )
        self.assertIsNone(
            fixer._detect_heading_level(
                "2.3.4.5 更深也归三级", bold=False, font_name="SimHei"
            )
        )
        self.assertIsNone(fixer._detect_heading_level("1.1 编制目的", bold=False))
        self.assertIsNone(
            fixer._detect_heading_level(
                "1.1.1 细则", bold=False, font_name="SimSun"
            )
        )
        # Not headings under pattern rules
        self.assertIsNone(fixer._detect_heading_level("（一）指挥长职责"))
        self.assertIsNone(fixer._detect_heading_level("（1）通过调查了解本厂区情况。"))
        self.assertIsNone(fixer._detect_heading_level("1 普通条目不应作标题"))
        self.assertIsNone(fixer._detect_heading_level("第一节 不适用"))
        self.assertIsNone(fixer._detect_heading_level("第一章总则................1"))
        self.assertIsNone(fixer._detect_heading_level("这是普通正文段落内容足够长。"))

    def test_numbering_conflict_rejects_out_of_order(self) -> None:
        """Reject numeric title that conflicts with established outline."""
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "conflict.docx"
            doc = Document()
            p = doc.add_paragraph()
            r = p.add_run("第一章 总则")
            r.bold = True
            r.font.size = Pt(14)
            p2 = doc.add_paragraph()
            r2 = p2.add_run("1.1 编制目的")
            r2.bold = True
            p3 = doc.add_paragraph()
            r3 = p3.add_run("1.2 适用范围")
            r3.bold = True
            # Body-like: bold but wrong chapter number → structural conflict
            p4 = doc.add_paragraph()
            r4 = p4.add_run("5.9 见前述条款")
            r4.bold = True
            doc.save(str(path))

            stats = StructureFixer().apply(path)
            self.assertGreaterEqual(stats["headings_promoted"], 3)
            self.assertGreaterEqual(stats.get("headings_rejected_conflict", 0), 1)

            doc2 = Document(str(path))
            texts = [(p.text or "").strip() for p in doc2.paragraphs]
            # 5.9 kept as body text (not stripped / not heading-only wording)
            self.assertTrue(any("5.9" in t for t in texts))
            headings = [
                (p.style.name if p.style else "", (p.text or "").strip())
                for p in doc2.paragraphs
                if (p.style and p.style.name or "").startswith("Heading")
            ]
            self.assertFalse(any("见前述" in t for _, t in headings))

    def test_detect_by_font_size_does_not_invent(self) -> None:
        """Font/size alone must NOT create headings without 第X章 / X.Y / X.Y.Z."""
        fixer = StructureFixer()
        self.assertIsNone(
            fixer._detect_heading_level(
                "环境突发事件应急预案",
                size_pt=18,
                bold=True,
                font_name="SimHei",
                body_median=12,
            )
        )
        self.assertIsNone(
            fixer._detect_heading_level(
                "编制说明",
                size_pt=14,
                bold=True,
                font_name="SimSun",
                body_median=12,
            )
        )

    def test_auto_number_keeps_title_font_size_bold(self) -> None:
        """Word owns 编号; title wording + font/size/bold preserved."""
        from docx.oxml.ns import qn
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "t.docx"
            doc = Document()
            p = doc.add_paragraph()
            r = p.add_run("第一章 总则")
            r.font.name = "SimSun"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
            r.font.size = Pt(14)
            r.bold = True
            doc.add_paragraph("普通正文")
            p2 = doc.add_paragraph()
            r2 = p2.add_run("1.1 编制目的")
            r2.font.name = "SimSun"
            r2._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
            r2.font.size = Pt(12)
            r2.bold = True
            doc.add_paragraph("（一）不应升为标题")
            doc.save(str(path))

            stats = StructureFixer().apply(path)
            self.assertGreaterEqual(stats["headings_promoted"], 2)
            self.assertGreaterEqual(stats["headings_numbered"], 2)
            self.assertEqual(stats["headings_promoted"], stats["headings_numbered"])

            doc2 = Document(str(path))
            texts = [(p.text or "").strip() for p in doc2.paragraphs if (p.text or "").strip()]
            self.assertIn("总则", texts)
            self.assertIn("编制目的", texts)
            self.assertIn("（一）不应升为标题", texts)
            self.assertNotIn("第一章 总则", texts)

            h1 = next(p for p in doc2.paragraphs if (p.text or "").strip() == "总则")
            self.assertEqual(h1.style.name, "Heading 1")
            self.assertIsNotNone(h1._p.get_or_add_pPr().find(qn("w:numPr")))
            run = next(r for r in h1.runs if (r.text or "").strip())
            r_fonts = run._element.find(qn("w:rPr")).find(qn("w:rFonts"))
            self.assertEqual(r_fonts.get(qn("w:eastAsia")), "SimSun")
            self.assertEqual(run.font.size, Pt(14))
            self.assertTrue(run.bold)

            paren = next(
                p for p in doc2.paragraphs if (p.text or "").strip() == "（一）不应升为标题"
            )
            self.assertFalse((paren.style.name or "").startswith("Heading"))

            style_num = doc2.styles["Heading 1"].element.get_or_add_pPr().find(qn("w:numPr"))
            self.assertIsNotNone(style_num)

            # List label rPr must match heading content (font/size/bold)
            num_pr = h1._p.get_or_add_pPr().find(qn("w:numPr"))
            used_num_id = int(num_pr.find(qn("w:numId")).get(qn("w:val")))
            numbering = doc2.part.numbering_part.element
            num_el = next(
                n
                for n in numbering.findall(qn("w:num"))
                if n.get(qn("w:numId")) == str(used_num_id)
            )
            override = num_el.find(qn("w:lvlOverride"))
            self.assertIsNotNone(override)
            lvl_rpr = override.find(qn("w:lvl")).find(qn("w:rPr"))
            self.assertIsNotNone(lvl_rpr)
            lvl_fonts = lvl_rpr.find(qn("w:rFonts"))
            self.assertEqual(lvl_fonts.get(qn("w:eastAsia")), "SimSun")
            self.assertEqual(lvl_rpr.find(qn("w:sz")).get(qn("w:val")), "28")  # 14pt
            b = lvl_rpr.find(qn("w:b"))
            self.assertIsNotNone(b)
            self.assertNotIn(b.get(qn("w:val")), ("0", "false", "off"))

    def test_partial_bold_number_only_is_emphasized(self) -> None:
        """pdf2docx often bolds only '1.1'; title text stays unbold — still heading."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "partial.docx"
            doc = Document()
            # Body sized 12pt for median
            body = doc.add_paragraph()
            br = body.add_run("这是一段足够长的正文内容用来估计正文字号中位数十二点。")
            br.font.size = Pt(12)
            # Chapter
            p0 = doc.add_paragraph()
            r0 = p0.add_run("第一章 总则")
            r0.font.size = Pt(14)
            # Partial bold: number bold, title not
            p = doc.add_paragraph()
            r_num = p.add_run("1.1")
            r_num.bold = True
            r_num.font.size = Pt(14)
            r_title = p.add_run("编制目的")
            r_title.bold = False
            r_title.font.size = Pt(14)
            # Explicit w:b val=0 on title (pdf2docx style)
            r_pr = r_title._element.get_or_add_rPr()
            b = r_pr.find(qn("w:b"))
            if b is None:
                b = OxmlElement("w:b")
                r_pr.append(b)
            b.set(qn("w:val"), "0")
            # Unemphasized body-like X.Y at 12pt, no bold → not heading
            p2 = doc.add_paragraph()
            r2 = p2.add_run("1.9 见前述说明条款内容")
            r2.bold = False
            r2.font.size = Pt(12)
            b2 = OxmlElement("w:b")
            b2.set(qn("w:val"), "0")
            r2._element.get_or_add_rPr().append(b2)
            doc.save(str(path))

            fixer = StructureFixer()
            doc_chk = Document(str(path))
            partial = next(
                p for p in doc_chk.paragraphs if (p.text or "").startswith("1.1")
            )
            self.assertTrue(fixer._leading_number_is_bold(partial))
            self.assertTrue(fixer._para_is_emphasized(partial, 12.0))
            # Majority alone would be 3/7 < 0.5 without leading-number rule
            bold_chars = sum(
                len((r.text or "").strip())
                for r in partial.runs
                if fixer._run_is_bold(r) and (r.text or "").strip()
            )
            total = sum(
                len((r.text or "").strip())
                for r in partial.runs
                if (r.text or "").strip()
            )
            self.assertLess(bold_chars / total, 0.5)

            stats = fixer.apply(path)
            self.assertGreaterEqual(stats["headings_promoted"], 2)
            doc2 = Document(str(path))
            headings = {
                (p.text or "").strip(): (p.style.name if p.style else "")
                for p in doc2.paragraphs
                if p.style and (p.style.name or "").startswith("Heading")
            }
            self.assertIn("编制目的", headings)
            self.assertTrue(headings["编制目的"].startswith("Heading"))
            # After strip, title wording should NOT inherit number-only bold
            h2 = next(p for p in doc2.paragraphs if (p.text or "").strip() == "编制目的")
            title_run = next(r for r in h2.runs if (r.text or "").strip())
            self.assertFalse(fixer._run_is_bold(title_run))

    def test_split_preserves_partial_bold_runs(self) -> None:
        """Packed title+body keeps number-bold / title-unbold as separate runs."""
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "packed_partial.docx"
            doc = Document()
            para = doc.add_paragraph()
            r1 = para.add_run("1.1")
            r1.bold = True
            r1.font.size = Pt(14)
            r2 = para.add_run("编制目的\n为有效应对突发环境事件，制定本预案。")
            r2.bold = False
            r2.font.size = Pt(12)
            doc.save(str(path))

            fixer = StructureFixer()
            stats = {"paragraphs_split": 0}
            d = Document(str(path))
            fixer._split_packed_paragraphs(d, stats)
            self.assertGreaterEqual(stats["paragraphs_split"], 1)
            title = next(p for p in d.paragraphs if (p.text or "").strip() == "1.1编制目的")
            self.assertTrue(fixer._leading_number_is_bold(title))
            runs = [(r.text, fixer._run_is_bold(r)) for r in title.runs if r.text]
            self.assertEqual(runs[0], ("1.1", True))
            self.assertFalse(runs[1][1])

    def test_body_left_indent_becomes_first_line(self) -> None:
        """pdf2docx left-only indent on body → 首行缩进; keep headings untouched."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "indent.docx"
            doc = Document()
            # Heading with left indent should stay (not body)
            h = doc.add_paragraph()
            hr = h.add_run("第一章 总则")
            hr.bold = True
            hr.font.size = Pt(14)
            # Body wrongly using whole-paragraph left indent (360 twips)
            body = doc.add_paragraph()
            br = body.add_run(
                "这是一段足够长的正文内容，应当使用首行缩进而不是整段左缩进。"
            )
            br.font.size = Pt(12)
            p_pr = body._p.get_or_add_pPr()
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), "360")
            ind.set(qn("w:right"), "144")
            p_pr.append(ind)
            # Body with both left + firstLine → drop left
            body2 = doc.add_paragraph()
            b2 = body2.add_run("另一段正文也有双倍缩进问题需要纠正成仅首行缩进。")
            b2.font.size = Pt(12)
            p_pr2 = body2._p.get_or_add_pPr()
            ind2 = OxmlElement("w:ind")
            ind2.set(qn("w:left"), "360")
            ind2.set(qn("w:firstLine"), "480")
            p_pr2.append(ind2)
            # Cover-like symmetric indent — leave alone
            cover = doc.add_paragraph("封面标题整段居中缩进示例文字")
            p_prc = cover._p.get_or_add_pPr()
            indc = OxmlElement("w:ind")
            indc.set(qn("w:left"), "1296")
            indc.set(qn("w:right"), "1296")
            p_prc.append(indc)
            doc.save(str(path))

            stats = StructureFixer().apply(path)
            self.assertGreaterEqual(stats.get("indents_normalized", 0), 2)

            doc2 = Document(str(path))
            paras = [p for p in doc2.paragraphs if (p.text or "").strip()]

            def read_ind(p):
                p_pr = p._p.find(qn("w:pPr"))
                if p_pr is None:
                    return {}
                el = p_pr.find(qn("w:ind"))
                if el is None:
                    return {}
                return {k.split("}")[-1]: el.get(k) for k in el.attrib}

            body_fixed = next(p for p in paras if "应当使用首行缩进" in (p.text or ""))
            info = read_ind(body_fixed)
            self.assertNotIn("left", info)
            self.assertEqual(info.get("firstLine"), "480")

            body2_fixed = next(p for p in paras if "双倍缩进" in (p.text or ""))
            info2 = read_ind(body2_fixed)
            self.assertNotIn("left", info2)
            self.assertEqual(info2.get("firstLine"), "480")

            cover_p = next(p for p in paras if "封面标题" in (p.text or ""))
            info_c = read_ind(cover_p)
            self.assertEqual(info_c.get("left"), "1296")
            self.assertEqual(info_c.get("right"), "1296")

    def test_demotes_paren_list_false_headings(self) -> None:
        """（1）列表项若被标成 Heading，必须降回正文。"""
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "paren.docx"
            doc = Document()
            p0 = doc.add_paragraph()
            r0 = p0.add_run("第一章 总则")
            r0.bold = True
            r0.font.size = Pt(14)
            p1 = doc.add_paragraph()
            r1 = p1.add_run("1.1 编制目的")
            r1.bold = True
            r1.font.size = Pt(14)
            # Simulate leftover false Heading from older conversion
            p2 = doc.add_paragraph("（1）预处理构筑物")
            p2.style = "Heading 3"
            p3 = doc.add_paragraph("（2）配水井")
            p3.style = "Heading 3"
            doc.add_paragraph("这是普通正文。")
            doc.save(str(path))

            stats = StructureFixer().apply(path)
            self.assertGreaterEqual(stats.get("headings_demoted", 0), 2)

            doc2 = Document(str(path))
            by_text = {
                (p.text or "").strip(): (p.style.name if p.style else "")
                for p in doc2.paragraphs
                if (p.text or "").strip()
            }
            self.assertFalse(by_text["（1）预处理构筑物"].startswith("Heading"))
            self.assertFalse(by_text["（2）配水井"].startswith("Heading"))
            self.assertTrue(by_text.get("总则", "").startswith("Heading"))
            self.assertTrue(by_text.get("编制目的", "").startswith("Heading"))

    def test_split_multiline_then_outline(self) -> None:
        """pdf2docx-style packed paragraph must become separate outline levels."""
        from docx.oxml.ns import qn
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "packed.docx"
            doc = Document()
            para = doc.add_paragraph()
            run = para.add_run(
                "第一章 总则\n1.1 编制目的\n1.1.1 细则\n为有效应对突发环境事件，制定本预案。"
            )
            run.font.name = "SimSun"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
            run.font.size = Pt(12)
            run.bold = True
            para2 = doc.add_paragraph()
            run2 = para2.add_run("第二章 组织机构\n2.1 应急指挥部")
            run2.font.name = "SimSun"
            run2._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
            run2.font.size = Pt(12)
            doc.save(str(path))

            stats = StructureFixer().apply(path)
            self.assertGreaterEqual(stats["paragraphs_split"], 4)

            doc2 = Document(str(path))
            by_text = {
                (p.text or "").strip(): p
                for p in doc2.paragraphs
                if (p.text or "").strip()
            }
            self.assertIn("总则", by_text)
            self.assertIn("编制目的", by_text)
            self.assertIn("细则", by_text)
            self.assertIn("为有效应对突发环境事件，制定本预案。", by_text)

            h1 = by_text["总则"]
            h2 = by_text["编制目的"]
            h3 = by_text["细则"]
            body = by_text["为有效应对突发环境事件，制定本预案。"]
            self.assertEqual(h1.style.name, "Heading 1")
            self.assertEqual(h2.style.name, "Heading 2")
            self.assertEqual(h3.style.name, "Heading 3")
            self.assertFalse((body.style.name or "").startswith("Heading"))
            self.assertIsNotNone(h1._p.get_or_add_pPr().find(qn("w:numPr")))

            run = next(r for r in h1.runs if (r.text or "").strip())
            self.assertTrue(run.bold)
            self.assertEqual(run.font.size, Pt(12))


class LLMClientParseTests(unittest.TestCase):
    def test_parse_json(self) -> None:
        raw = '{"items":[{"index":0,"style":"Heading 1"},{"index":1,"style":"Body"}]}'
        items = LLMClient._parse_items(raw)
        self.assertEqual(len(items), 2)
        self.assertEqual(items[0]["style"], "Heading 1")

    def test_parse_fenced(self) -> None:
        raw = '```json\n{"items":[{"index":3,"style":"Heading 2"}]}\n```'
        items = LLMClient._parse_items(raw)
        self.assertEqual(items[0]["index"], 3)


if __name__ == "__main__":
    unittest.main()
